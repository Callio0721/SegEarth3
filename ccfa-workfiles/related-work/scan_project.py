import ast, hashlib, json, os, re, sys
from pathlib import Path

ROOT = Path(r"D:\论文\9.27论文")
OUT = ROOT / "ccfa-workfiles" / "related-work" / "project-inventory.json"
IGNORE_DIRS = {".git", ".obsidian", "__pycache__", ".pytest_cache", ".mypy_cache"}
TEXT_SUFFIXES = {".md", ".txt", ".py", ".sh", ".yaml", ".yml", ".json", ".csv", ".bib", ".html", ".xml", ".toml", ".ini", ".cfg"}
SUFFIXES = {}
rows = []
errors = []

def digest(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()

def safe_text(path, limit=None):
    try:
        data = path.read_text(encoding="utf-8-sig", errors="ignore")
    except Exception as exc:
        errors.append({"path": str(path), "stage": "text", "error": repr(exc)})
        return ""
    return data if limit is None else data[:limit]

def summarize_text(path):
    text = safe_text(path)
    return {
        "lines": text.count("\n") + (1 if text else 0),
        "chars": len(text),
        "headings": [line.strip() for line in text.splitlines() if re.match(r"^\s*#{1,6}\s+", line)][:40],
        "keywords": {k: text.lower().count(k) for k in ["segmentation", "open-vocabulary", "prompt", "sam3", "sam 3", "clip", "grounding", "context", "remote sensing"]},
        "excerpt": text[:1200],
    }

def summarize_python(path):
    text = safe_text(path)
    info = {"lines": text.count("\n") + (1 if text else 0), "chars": len(text), "classes": [], "functions": [], "imports": []}
    try:
        tree = ast.parse(text)
        for node in tree.body:
            if isinstance(node, ast.ClassDef):
                info["classes"].append({"name": node.name, "line": node.lineno})
            elif isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
                info["functions"].append({"name": node.name, "line": node.lineno})
            elif isinstance(node, ast.Import):
                info["imports"].extend(alias.name for alias in node.names)
            elif isinstance(node, ast.ImportFrom):
                info["imports"].append(node.module or "")
    except Exception as exc:
        info["syntax_error"] = repr(exc)
    return info

def pdf_summary(path):
    try:
        import fitz
        doc = fitz.open(str(path))
        chunks = []
        for page in doc:
            chunks.append(page.get_text("text"))
        text = "\n".join(chunks)
        first = " ".join(text.split())[:500]
        lower = text.lower()
        terms = ["open-vocabulary", "open vocabulary", "segmentation", "remote sensing", "sam 3", "sam3", "clip", "prompt", "grounding", "context", "global", "local", "concept"]
        info = {
            "pages": len(doc),
            "chars": len(text),
            "metadata_title": (doc.metadata or {}).get("title", ""),
            "first_text": first,
            "terms": {term: lower.count(term) for term in terms},
        }
        doc.close()
        return info
    except Exception as exc:
        errors.append({"path": str(path), "stage": "pdf", "error": repr(exc)})
        return {"pages": None, "chars": None, "error": repr(exc)}

def docx_summary(path):
    try:
        import docx
        doc = docx.Document(str(path))
        parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(parts)
        return {"paragraphs": len(parts), "chars": len(text), "excerpt": text[:1500]}
    except Exception as exc:
        errors.append({"path": str(path), "stage": "docx", "error": repr(exc)})
        return {"error": repr(exc)}

def pptx_summary(path):
    try:
        from pptx import Presentation
        prs = Presentation(str(path))
        parts = []
        for idx, slide in enumerate(prs.slides, 1):
            slide_parts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_parts.append(shape.text.strip())
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        slide_parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            if slide_parts:
                parts.append(f"SLIDE {idx}: " + "\n".join(slide_parts))
        text = "\n\n".join(parts)
        return {"slides": len(prs.slides), "chars": len(text), "excerpt": text[:1800]}
    except Exception as exc:
        errors.append({"path": str(path), "stage": "pptx", "error": repr(exc)})
        return {"error": repr(exc)}

for path in ROOT.rglob("*"):
    if not path.is_file():
        continue
    if any(part in IGNORE_DIRS for part in path.parts):
        continue
    if path.name.startswith("._"):
        continue
    rel = str(path.relative_to(ROOT))
    suffix = path.suffix.lower()
    SUFFIXES[suffix] = SUFFIXES.get(suffix, 0) + 1
    row = {"path": rel, "suffix": suffix, "bytes": path.stat().st_size, "sha256": digest(path)}
    try:
        if suffix == ".py":
            row.update(summarize_python(path))
        elif suffix in TEXT_SUFFIXES:
            row.update(summarize_text(path))
        elif suffix == ".pdf":
            row.update(pdf_summary(path))
        elif suffix == ".docx":
            row.update(docx_summary(path))
        elif suffix == ".pptx":
            row.update(pptx_summary(path))
        elif suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".svg"}:
            if suffix != ".svg":
                try:
                    from PIL import Image
                    with Image.open(path) as img:
                        row["image_size"] = [img.width, img.height]
                except Exception as exc:
                    row["image_error"] = repr(exc)
    except Exception as exc:
        errors.append({"path": str(path), "stage": "outer", "error": repr(exc)})
    rows.append(row)

result = {
    "root": str(ROOT),
    "file_count": len(rows),
    "suffix_counts": dict(sorted(SUFFIXES.items(), key=lambda item: (-item[1], item[0]))),
    "errors": errors,
    "files": rows,
}
OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps({"output": str(OUT), "file_count": len(rows), "suffix_counts": result["suffix_counts"], "errors": len(errors)}, ensure_ascii=False, indent=2))
